# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "msoffcrypto-tool==5.4.2",
#   "odfpy==1.4.1",
#   "openpyxl==3.1.5",
#   "xlwt==1.3.0",
# ]
# ///
"""Regenerate the spreadsheet fixtures: `uv run generate.py` from this directory.

Every format holds the same `Summary` sheet, so the Rust tests and the wrapper's
integration tests can assert one set of expectations against all of them. Written by
the libraries that are each format's reference writer rather than assembled by hand:
BIFF8 inside a compound file is not something to hand-roll in a test.

None of these writers computes a formula's cached result, and a reader only ever sees
that cached result, so there are deliberately no formulas here.
"""

import datetime as dt
from pathlib import Path

from msoffcrypto.format.ooxml import OOXMLFile
import openpyxl
import xlwt
from odf.opendocument import OpenDocumentSpreadsheet
from odf.table import Table, TableCell, TableRow
from odf.text import P

HERE = Path(__file__).parent

ROWS = [
    ["Item", "Qty", "Price", "Due"],
    ["Bolts", 12, 0.25, dt.date(2024, 3, 15)],
    # An em dash: outside Latin-1, and the character PDFDocEncoding gets wrong too.
    ["Nuts — M8", 200, 0.1, dt.datetime(2024, 3, 15, 13, 30)],
    ["Total", None, 23, None],
    [True, False, None, None],
]


def xlsx() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in ROWS:
        ws.append(row)
    ws.merge_cells("A4:B4")
    ws["C5"] = "#DIV/0!"  # openpyxl stores a known error string as an error cell

    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["F10"] = "far"
    wb.save(HERE / "sample.xlsx")


def styled() -> None:
    """Number formats, fonts, fills, borders, alignment and layout, one of each."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "Quarterly report"
    ws["A1"].font = Font(bold=True, size=22, color="FF1F4E79")
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 30

    header = ["Item", "Amount", "Share", "Due"]
    thin = Side(style="thin", color="FF000000")
    for c, text in enumerate(header, start=1):
        cell = ws.cell(row=2, column=c, value=text)
        cell.font = Font(bold=True, color="FFFFFFFF")
        cell.fill = PatternFill("solid", fgColor="FF4472C4")
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=Side(style="double", color="FF0000FF"))

    rows = [
        ("Steel", 1234.5, 0.256, dt.date(2024, 3, 15)),
        ("Concrete", -987.25, 0.5, dt.date(2024, 12, 1)),
    ]
    for r, (item, amount, share, due) in enumerate(rows, start=3):
        ws.cell(row=r, column=1, value=item).font = Font(italic=True)
        a = ws.cell(row=r, column=2, value=amount)
        a.number_format = '#,##0.00;[Red](#,##0.00)'
        s = ws.cell(row=r, column=3, value=share)
        s.number_format = "0.0%"
        d = ws.cell(row=r, column=4, value=due)
        d.number_format = "d-mmm-yy"
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = Border(left=thin, right=thin, top=thin, bottom=thin)

    note = ws.cell(row=5, column=1, value="A note long enough that it has to wrap")
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=5, column=5).fill = PatternFill("solid", fgColor="FFFFFF00")

    ws.column_dimensions["A"].width = 20.7109375
    ws.column_dimensions["F"].hidden = True
    ws.row_dimensions[6].hidden = True
    wb.save(HERE / "styled.xlsx")


def extent() -> None:
    """Two cells at opposite corners of the sheet: tiny file, enormous bounding box."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "top"
    ws["XFD1048576"] = "bottom"
    wb.save(HERE / "extent.xlsx")


def encrypted() -> None:
    # Agile encryption, Excel's default since 2010. The result is a compound file, not
    # a zip, which is what makes it look like an `.xls` to a sniffer.
    with open(HERE / "sample.xlsx", "rb") as plain, open(
        HERE / "encrypted.xlsx", "wb"
    ) as out:
        office = OOXMLFile(plain)
        office.encrypt("papyra", out)


def xls() -> None:
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Summary")
    date = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    stamp = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM")
    for r, row in enumerate(ROWS):
        for c, value in enumerate(row):
            if value is None or (r == 3 and c < 2):
                continue
            if isinstance(value, dt.datetime):
                ws.write(r, c, value, stamp)
            elif isinstance(value, dt.date):
                ws.write(r, c, value, date)
            else:
                ws.write(r, c, value)
    ws.write_merge(3, 3, 0, 1, "Total")

    hidden = wb.add_sheet("Hidden")
    hidden.visibility = 1
    hidden.write(9, 5, "far")
    wb.save(str(HERE / "sample.xls"))


def ods() -> None:
    doc = OpenDocumentSpreadsheet()
    table = Table(name="Summary")
    for row in ROWS:
        tr = TableRow()
        for value in row:
            if value is None:
                tr.addElement(TableCell())
            elif isinstance(value, bool):
                tr.addElement(
                    TableCell(valuetype="boolean", booleanvalue=str(value).lower())
                )
            elif isinstance(value, (int, float)):
                cell = TableCell(valuetype="float", value=value)
                cell.addElement(P(text=str(value)))
                tr.addElement(cell)
            elif isinstance(value, dt.date):
                cell = TableCell(valuetype="date", datevalue=value.isoformat())
                cell.addElement(P(text=value.isoformat()))
                tr.addElement(cell)
            else:
                cell = TableCell(valuetype="string")
                cell.addElement(P(text=value))
                tr.addElement(cell)
        table.addElement(tr)
    doc.spreadsheet.addElement(table)
    doc.save(str(HERE / "sample.ods"))


if __name__ == "__main__":
    xlsx()
    styled()
    extent()
    encrypted()
    xls()
    ods()
